import streamlit as st
from pypdf import PdfReader
from pptx import Presentation
import json
from google import genai
from google.genai import types

# 1. API Ayarı
client = genai.Client(api_key="AQ.Ab8RN6LiJFs1kzywfqPeAMBs65RZsj-6UlqSgjtO13E8Sgk62Q")

st.set_page_config(page_title="SlideQuiz", page_icon="📝", layout="centered")
st.title("SlideQuiz: Slayttan Sınav Oluşturucu")

# Session state başlatma
if "questions" not in st.session_state:
    st.session_state.questions = []
if "user_answers" not in st.session_state:
    st.session_state.user_answers = {}
if "submitted" not in st.session_state:
    st.session_state.submitted = False

# Dosyadan metin ayıklama fonksiyonu
def extract_text(file):
    text = ""
    if file.name.endswith(".pdf"):
        reader = PdfReader(file)
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text += t + "\n"
    elif file.name.endswith(".pptx"):
        prs = Presentation(file)
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text += shape.text + "\n"
    return text

# Dosya Yükleme Ekranı
uploaded_file = st.file_uploader("Ders notunu veya sunumunu yükle (PDF / PPTX)", type=["pdf", "pptx"])

if uploaded_file and not st.session_state.questions:
    q_count = st.slider("Kaç soru üretilsin?", min_value=3, max_value=10, value=5)
    
    if st.button("Sınavı Oluştur"):
        with st.spinner("Materyal taranıyor ve sorular hazırlanıyor..."):
            extracted_text = extract_text(uploaded_file)
            
            prompt = f"""
            Aşağıdaki metne dayanarak {q_count} adet çoktan seçmeli soru hazırla.
            Her sorunun 4 seçeneği (A, B, C, D) ve 1 doğru cevabı olsun.
            
            İçerik:
            {extracted_text[:4000]}
            """
            
            # JSON formatında çıktı alma
            response = client.models.generate_content(
                model="gemini-2.5-flash",
                contents=prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    response_schema={
                        "type": "ARRAY",
                        "items": {
                            "type": "OBJECT",
                            "properties": {
                                "question": {"type": "STRING"},
                                "options": {
                                    "type": "ARRAY",
                                    "items": {"type": "STRING"}
                                },
                                "answer": {"type": "STRING"}
                            },
                            "required": ["question", "options", "answer"]
                        }
                    }
                )
            )
            st.session_state.questions = json.loads(response.text)
            st.rerun()

# Sınav Ekranı
if st.session_state.questions and not st.session_state.submitted:
    st.subheader("Sınav")
    for idx, q in enumerate(st.session_state.questions):
        st.write(f"**Soru {idx + 1}:** {q['question']}")
        st.session_state.user_answers[idx] = st.radio(
            "Seçiminiz:",
            options=q["options"],
            key=f"q_{idx}",
            index=None
        )
        st.write("---")
        
    if st.button("Sınavı Tamamla"):
        st.session_state.submitted = True
        st.rerun()

# Sonuç Ekranı
if st.session_state.submitted:
    st.subheader("Sınav Sonuçları")
    correct_count = 0
    total = len(st.session_state.questions)
    
    for idx, q in enumerate(st.session_state.questions):
        user_choice = st.session_state.user_answers.get(idx)
        is_correct = user_choice == q["answer"]
        if is_correct:
            correct_count += 1
            st.success(f"Soru {idx + 1}: Doğru!")
        else:
            st.error(f"Soru {idx + 1}: Yanlış! (Senin cevabın: {user_choice} | Doğru: {q['answer']})")
            
    score = int((correct_count / total) * 100)
    st.metric(label="Başarı Puanı", value=f"%{score} ({correct_count}/{total})")
    
    if st.button("Yeni Sınav Oluştur"):
        st.session_state.questions = []
        st.session_state.user_answers = {}
        st.session_state.submitted = False
        st.rerun()